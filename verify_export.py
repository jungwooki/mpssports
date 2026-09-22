# -*- coding: utf-8 -*-
from pathlib import Path
from playwright.sync_api import sync_playwright
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches
import json,re,sys
D=Path(__file__).resolve().parent
REF=D.parent/(D.name+'_1');EXPORT=REF/'내보내기';EXPORT.mkdir(parents=True,exist_ok=True)
PRET='--pretendard' in sys.argv
FONT='Pretendard' if PRET else 'A2Z'
SUFFIX='_Pretendard' if PRET else ''
HTML='index_pretendard.html' if PRET else 'index.html'
out=REF/('preview_pretendard' if PRET else 'preview');out.mkdir(exist_ok=True)
with sync_playwright() as w:
 b=w.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
 p=b.new_page(viewport={'width':1600,'height':962},device_scale_factor=1)
 errors=[];p.on('pageerror',lambda e:errors.append(str(e)))
 p.goto((D/HTML).as_uri());p.evaluate('document.fonts.ready');p.wait_for_timeout(300)
 n=p.locator('.slide').count();issues=[];title_issues=[]
 for i in range(n):
  p.evaluate('(i)=>show(i)',i)
  p.locator('.slide.active').screenshot(path=str(out/f'{i+1:02}.png'))
  bad=p.evaluate('''()=>{let s=document.querySelector('.slide.active'),sr=s.getBoundingClientRect(),f=s.querySelector('footer').getBoundingClientRect();return [...s.querySelectorAll('.content h2,.content h3,.content p,.content .band,.content table,.content .stat,.content .flow,.content .team-support,.content .source-grid,.content .warm-photo,.content .dropout-aside,.content .reference-toc,.content .service-tiles,.content .network-logos,.content .unit-summary,.content .report-story,.content .report-reading,.content .report-service,.content .report-payment,.heading')].flatMap(e=>{let r=e.getBoundingClientRect();return r.bottom>f.top-6||r.right>sr.right-45?[{text:e.innerText.slice(0,90),bottom:Math.round(r.bottom),footer:Math.round(f.top),right:Math.round(r.right)}]:[]})}''')
  if bad:issues.append({'slide':i+1,'issues':bad})
  title=p.evaluate('''()=>{let e=document.querySelector('.slide.active .heading h1'),s=document.querySelector('.slide.active'),r=document.createRange();r.selectNodeContents(e);let rects=[...r.getClientRects()].filter(x=>x.width>0),tops=new Set(rects.map(x=>Math.round(x.top)));return {lines:tops.size,overflow:e.scrollWidth>e.clientWidth+1||rects.some(x=>x.right>s.getBoundingClientRect().right-45),text:e.innerText}}''')
  if title['lines']!=1 or title['overflow']:title_issues.append({'slide':i+1,**title})
 broken=p.evaluate("[...document.images].filter(i=>i.getAttribute('src')&&(!i.complete||i.naturalWidth===0)).map(i=>i.alt)")
 p.evaluate('show(0)');p.locator('#next').click();assert p.locator('#counter').inner_text()==f'02 / {n}'
 p.keyboard.press('ArrowRight');assert p.locator('#counter').inner_text()==f'03 / {n}'
 p.locator('#contents').click();assert p.locator('.toc-grid button').count()==n;p.locator('.toc-grid button').nth(6).click();assert p.locator('#counter').inner_text()==f'07 / {n}'
 p.evaluate('show(7)');p.locator('.slide.active .zoom').first.click();assert p.locator('#image-dialog').evaluate('(e)=>e.open');p.keyboard.press('Escape')
 p.locator('#overview').click();assert p.locator('.slide:visible').count()==n;p.locator('#overview').click()
 p.set_viewport_size({'width':390,'height':844});p.evaluate('show(0)');p.screenshot(path=str(out/'mobile.png'));assert p.evaluate('document.documentElement.scrollWidth')==390
 p.set_viewport_size({'width':1600,'height':962});p.emulate_media(media='print')
 p.pdf(path=str(EXPORT/f'SPORTS_MPS_투자제안서_20260922{SUFFIX}.pdf'),width='1600px',height='900px',print_background=True,prefer_css_page_size=True)
 result={'font':FONT,'slides':n,'overflow':issues,'broken_images':broken,'js_errors':errors,'font_loaded':p.evaluate("font=>document.fonts.check('800 46px '+font)",FONT),'navigation':'pass','toc':'pass','report_zoom':'pass','overview':'pass','mobile_width':'pass','print':'pass'}
 result['single_line_title_issues']=title_issues
 (out/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False));b.close()
board=Image.new('RGB',(1600,((n+3)//4)*245),'#e9eaed');d=ImageDraw.Draw(board)
for i in range(n):
 im=Image.open(out/f'{i+1:02}.png');im.thumbnail((390,220));x=i%4*400+5;y=i//4*245+20;board.paste(im,(x,y));d.text((x,y-16),str(i+1),fill='black')
board.save(out/'전체미리보기.jpg')
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
notes=json.loads((REF/'참고문서/slide_notes.json').read_text())
for i in range(n):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.shapes.add_picture(str(out/f'{i+1:02}.png'),0,0,width=prs.slide_width,height=prs.slide_height)
 r=notes[i];s.notes_slide.notes_text_frame.text=re.sub('<[^>]+>',' ',r['title']+'\n'+r['sub']+'\n'+r['source']+'\n'+r['notes'])+'\n원본 편집: 동봉 index.html 및 build_deck.py. 본 PPTX는 레이아웃 보존을 위한 슬라이드 이미지형입니다.'
prs.save(EXPORT/f'SPORTS_MPS_투자제안서_20260922{SUFFIX}_발표용.pptx')
print('PDF and presentation PPTX exported')
